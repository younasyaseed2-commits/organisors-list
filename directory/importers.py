import csv
from io import TextIOWrapper

from django.db import transaction
from openpyxl import load_workbook

from .district_resolver import resolve_district_name
from .models import District, Location, Organizer


REQUIRED_COLUMNS = {"organizer"}
PHONE_COLUMNS = {"phone", "phone_number", "mobile", "mobile_number"}
LOCATION_COLUMNS = {"location", "locations"}
ALIAS_COLUMNS = {"alias", "aliases", "malayalam", "malayalam_name"}


def _normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _split_locations(value):
    if value is None:
        return []
    return [name.strip() for name in str(value).split(",") if name.strip()]


def _read_csv(upload_file):
    wrapper = TextIOWrapper(upload_file.file, encoding="utf-8-sig")
    reader = csv.DictReader(wrapper)
    for row in reader:
        yield {_normalize_header(key): value for key, value in row.items()}


def _read_xlsx(upload_file):
    workbook = load_workbook(upload_file, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [_normalize_header(value) for value in next(rows, [])]

    for row in rows:
        yield {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }


def _read_rows(upload_file):
    if upload_file.name.lower().endswith(".csv"):
        return _read_csv(upload_file)
    return _read_xlsx(upload_file)


def import_directory_file(upload_file):
    created = {
        "districts": 0,
        "organizers": 0,
        "locations": 0,
        "skipped_rows": 0,
    }

    rows = list(_read_rows(upload_file))
    if not rows:
        return created

    headers = set(rows[0].keys())
    if not REQUIRED_COLUMNS.issubset(headers):
        missing = ", ".join(sorted(REQUIRED_COLUMNS - headers))
        raise ValueError(f"Missing required column(s): {missing}")

    phone_column = next((column for column in PHONE_COLUMNS if column in headers), None)
    location_column = next((column for column in LOCATION_COLUMNS if column in headers), None)
    alias_column = next((column for column in ALIAS_COLUMNS if column in headers), None)

    if not phone_column:
        raise ValueError("Missing required phone column. Use phone or phone_number.")
    if not location_column:
        raise ValueError("Missing required location column. Use location or locations.")

    with transaction.atomic():
        for row in rows:
            organizer_name = str(row.get("organizer") or "").strip()
            phone_number = str(row.get(phone_column) or "").strip()
            locations = _split_locations(row.get(location_column))
            aliases = _split_locations(row.get(alias_column)) if alias_column else []

            if not organizer_name or not phone_number or not locations:
                created["skipped_rows"] += 1
                continue

            organizer_district_name = resolve_district_name(locations[0], row.get("district"))
            district, district_created = District.objects.get_or_create(
                name=organizer_district_name
            )
            organizer, organizer_created = Organizer.objects.get_or_create(
                name=organizer_name,
                phone_number=phone_number,
                district=district,
            )

            created["districts"] += int(district_created)
            created["organizers"] += int(organizer_created)

            for index, location_name in enumerate(locations):
                location_district_name = resolve_district_name(location_name, row.get("district"))
                location_district, location_district_created = District.objects.get_or_create(
                    name=location_district_name
                )
                alias = aliases[index] if index < len(aliases) else ""
                _, location_created = Location.objects.get_or_create(
                    name=location_name,
                    organizer=organizer,
                    defaults={"district": location_district, "aliases": alias},
                )
                created["districts"] += int(location_district_created)
                created["locations"] += int(location_created)

    return created
